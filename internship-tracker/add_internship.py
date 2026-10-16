#!/usr/bin/env python3
"""Add internships to Internship_Tracker.xlsx from the terminal.

Usage:
    python3 add_internship.py [path/to/Internship_Tracker.xlsx] [--details]

Type a company name and its role (Enter to skip the role); it's saved into the next empty row.
Press Enter on an empty line (or type q) to quit.
"""
import argparse
import datetime
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is missing. Install it with:  python3 -m pip install openpyxl")

FIRST_ROW, LAST_ROW = 12, 500
COMPANY, ROLE, LOCATION, LINK, STATUS, APPLIED, DEADLINE, DAYS = range(1, 9)
NOTES = 15
STATUSES = ["Wishlist", "Networking", "Applied", "OA/Assessment", "Interviewing",
            "Offer", "Accepted", "Rejected", "Withdrawn"]


def cell_text(ws, row, col):
    v = ws.cell(row, col).value
    return "" if v is None else str(v).strip()


def existing_companies(ws):
    return {cell_text(ws, r, COMPANY).lower(): r
            for r in range(FIRST_ROW, LAST_ROW + 1) if cell_text(ws, r, COMPANY)}


def next_empty_row(ws):
    for r in range(FIRST_ROW, LAST_ROW + 1):
        if not cell_text(ws, r, COMPANY):
            return r
    return None


def parse_date(s):
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m/%d"):
        try:
            d = datetime.datetime.strptime(s, fmt)
            if fmt == "%m/%d":
                d = d.replace(year=datetime.date.today().year)
            return d
        except ValueError:
            pass
    raise ValueError(f"Couldn't read date {s!r}; use MM/DD/YYYY.")


def ask(prompt, default=""):
    shown = f"{prompt} [{default}]: " if default else f"{prompt}: "
    answer = input(shown).strip()
    return answer or default


def ask_details():
    details = {
        "role": ask("  Role title"),
        "location": ask("  Location"),
    }
    while True:
        status = ask("  Status (" + ", ".join(STATUSES) + ")", "Applied")
        match = next((s for s in STATUSES if s.lower() == status.lower()), None)
        if match:
            details["status"] = match
            break
        print("  Not a valid status, try again.")
    for key, label in (("applied", "Date applied (MM/DD/YYYY)"), ("deadline", "Deadline (MM/DD/YYYY)")):
        default = datetime.date.today().strftime("%m/%d/%Y") if key == "applied" else ""
        while True:
            raw = ask("  " + label, default)
            try:
                details[key] = parse_date(raw) if raw else None
                break
            except ValueError as e:
                print("  " + str(e))
    details["link"] = ask("  Job link")
    details["notes"] = ask("  Notes")
    return details


def add_row(ws, row, company, d):
    ws.cell(row, COMPANY).value = company
    if d.get("role"):
        ws.cell(row, ROLE).value = d["role"]
    if d.get("location"):
        ws.cell(row, LOCATION).value = d["location"]
    if d.get("link"):
        ws.cell(row, LINK).value = d["link"]
        ws.cell(row, LINK).hyperlink = d["link"]
    ws.cell(row, STATUS).value = d.get("status", "Applied")
    if d.get("applied"):
        ws.cell(row, APPLIED).value = d["applied"]
    if d.get("deadline"):
        ws.cell(row, DEADLINE).value = d["deadline"]
    if d.get("notes"):
        ws.cell(row, NOTES).value = d["notes"]
    if not ws.cell(row, DAYS).value:
        ws.cell(row, DAYS).value = f'=IF(F{row}="","",TODAY()-F{row})'


def save(wb, path):
    while True:
        try:
            wb.save(path)
            return True
        except PermissionError:
            again = input("Can't save: the file is open in Excel. Close it, then press Enter to retry (or type q to skip): ")
            if again.strip().lower() == "q":
                return False


def main():
    parser = argparse.ArgumentParser(description="Quickly add internships to your tracker spreadsheet.")
    parser.add_argument("file", nargs="?", default="Internship_Tracker.xlsx", help="path to the tracker .xlsx")
    parser.add_argument("-d", "--details", action="store_true",
                        help="also ask for role, location, status, dates, link and notes")
    args = parser.parse_args()

    try:
        wb = openpyxl.load_workbook(args.file)
    except FileNotFoundError:
        sys.exit(f"Can't find {args.file}. Run this from the folder that has the file, "
                 f"or give the path:  python3 add_internship.py ~/Downloads/Internship_Tracker.xlsx")
    ws = wb["Tracker"] if "Tracker" in wb.sheetnames else wb.active

    print(f"Adding to {args.file}. Type a company and press Enter. Empty line or q to quit.\n")
    added = 0
    while True:
        try:
            company = input("Company: ").strip()
        except (EOFError, KeyboardInterrupt):
            print()
            break
        if not company or company.lower() == "q":
            break

        dup_row = existing_companies(ws).get(company.lower())
        if dup_row and ask(f"  {company} is already in row {dup_row}. Add anyway? (y/N)", "n").lower() != "y":
            continue

        row = next_empty_row(ws)
        if row is None:
            print(f"The tracker is full (rows {FIRST_ROW}-{LAST_ROW}).")
            break

        if args.details:
            details = ask_details()
        else:
            details = {"role": ask("  Role title"),
                       "applied": datetime.datetime.combine(datetime.date.today(), datetime.time())}
        add_row(ws, row, company, details)
        if save(wb, args.file):
            added += 1
            print(f"  ✓ Saved {company} to row {row}")

    print(f"Done. Added {added} {'company' if added == 1 else 'companies'}.")


if __name__ == "__main__":
    main()
